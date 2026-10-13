import os, copy
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

V3 = r"C:\Users\Kartavya\Desktop\NYU\Summer\Risk Premia\Analysis\Risk Premia Analysis v3_2012.docx"
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "v5img")


def new_doc():
    d = docx.Document(V3)
    body = d.element.body
    for el in list(body):
        if el.tag.endswith("}sectPr"):
            continue
        body.remove(el)
    return d


def _par(d, text, style="Normal", size=None, italic=None, bold=None, align=None):
    p = d.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        if size: r.font.size = Pt(size)
        if italic: r.italic = True
        if bold: r.bold = True
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def render(d, blocks):
    for b in blocks:
        t = b[0]
        if t == "h1": _par(d, b[1], "Heading 1")
        elif t == "h2": _par(d, b[1], "Heading 2")
        elif t == "h3": _par(d, b[1], "Heading 3")
        elif t == "sub": _par(d, b[1], "Normal", size=9, italic=True)
        elif t == "b": _par(d, b[1], "List Bullet")
        elif t == "p": _par(d, b[1], "Normal")
        elif t == "cap":                      # table caption
            _par(d, b[1], "Normal", size=9, bold=True)
        elif t == "fig":
            p = d.add_paragraph(style="Normal"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(IMG, b[1]), width=Inches(6.0))
            _par(d, b[2], "Normal", size=9, bold=True, align="center")
        elif t == "tbl":
            header, rows = b[1], b[2]
            tb = d.add_table(rows=1 + len(rows), cols=len(header))
            tb.style = d.styles["Light Grid Accent 1"]
            for j, h in enumerate(header):
                c = tb.rows[0].cells[j]; c.text = ""
                r = c.paragraphs[0].add_run(h); r.font.size = Pt(9)
            for i, row in enumerate(rows, 1):
                for j, v in enumerate(row):
                    c = tb.rows[i].cells[j]; c.text = ""
                    r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(9)
            d.add_paragraph("")
        elif t == "pb":
            d.add_page_break()
        else:
            raise ValueError(t)
