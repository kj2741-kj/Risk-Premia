"""Updates Analysis/Research_Methodology_v4.docx (12 October 2026, second revision). Additive and dated; Dimil's text is kept except the
'Statistical status' bullet of Section 14, which is rewritten. Run once from any folder; writes the file in place after a successful build."""
import copy, os, sys, tempfile
import docx
from docx.shared import Pt, Emu
from docx.oxml.ns import qn

PATH = r"C:\Users\Kartavya\Desktop\NYU\Summer\Risk Premia\Analysis\Research_Methodology_v4.docx"
d = docx.Document(PATH)
if any(p.text.startswith("5.9 Design Decisions") for p in d.paragraphs):
    sys.exit("already updated")
body = d.element.body


def find(prefix, style=None):
    hits = [p for p in d.paragraphs if p.text.startswith(prefix) and (style is None or p.style.name == style)]
    assert len(hits) == 1, (prefix, len(hits))
    return hits[0]


def after(anchor_el, new_el):
    anchor_el.addnext(new_el)
    return new_el


def mk_par(text="", style="Normal", lead=None, size=None, italic=False):
    p = d.add_paragraph(style=style)
    if lead:
        r = p.add_run(lead); r.bold = True
    if text:
        r = p.add_run(text)
        if size: r.font.size = Pt(size)
        if italic: r.italic = True
    body.remove(p._element)
    return p


def mk_table(header, rows, widths_in):
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t.style = d.styles["Table Grid"]
    t.autofit = False
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(9)
    for row in t.rows:
        for j, w in enumerate(widths_in):
            row.cells[j].width = Emu(int(w * 914400))
    body.remove(t._element)
    return t


# ---- revision note under "As of 12 October 2026"
asof = find("As of 12 October 2026")
note = mk_par("Revised 12 October 2026 (second revision): Section 5.9 (design decisions and their recorded basis), the Skewness implementation notes in Sections 5.8, 12, 14 and 15, "
              "and Section 16 (Streamlit dashboards and the hub) are new. The Statistical status bullet of Section 14 is rewritten and two bullets are added to it.", size=11, italic=True)
after(asof._element, note._element)

# ---- Section 1: pointer to Section 16
s1 = find("The dashboard combines sleeves and asset classes by two methods only")
s1.add_run(" Section 16 describes how the Streamlit dashboards and the hub implement the same sleeves.")

# ---- Section 5.8: implementation notes
s58 = find("For each product, the signal is minus the skewness")
impl = mk_par(
    "Implementation notes (added 12 October 2026). (a) The original skewness factor is cross-sectional: Fernandez-Perez, Frijns, Fuertes and Miffre (2018, Journal of Banking and Finance 86, 143-158) go long the most negatively skewed and short the most positively skewed commodities; the 2021 paper applies it to energy futures. "
    "(b) Warm-up: a full 252-observation window of continuous-series returns is required, so a product's first position starts two trading days after its first signal, between 28 December 2006 and 9 January 2007 for Metals, Energy, Precious Metals and Agriculture and between 19 August and 15 December 2010 for NGL (Isobutane is the latest). Every product is therefore live before the reporting window opens. "
    "(c) A missing return inside the window would blank the signal until the window is full again; the continuous series carries a zero return over a missing price (Section 4.2) and has no missing returns in the data, so this does not occur. "
    "(d) The sleeve is not market neutral: it uses one product's own trailing year and can keep the same sign for long periods, for example while a large collapse stays inside the window or when repeated shocks keep the skewness negative. "
    "(e) The signal is computed on the ratio-stitched continuous series of Section 4.1; the Streamlit dashboards rebuild the same returns from their additive series (Section 16).")
after(s58._element, impl._element)

# ---- Section 5.9: design decisions
h59 = mk_par("5.9 Design Decisions and Their Recorded Basis", style="Heading 3")
after(impl._element, h59._element)
intro59 = mk_par("Every parameter above was set by a research decision recorded in ENGINE_CHANGELOG.md or in Readme Portfolio construction 10 Oct.md; none is estimated or optimized by the dashboard code. The aim of the decisions is to keep the rules close to implementable practice. "
                 "Where an alternative was compared, the comparison was run on the same 2011 to 2026 sample, on full-sample IR, and was not tested for significance, so the table records what each decision rests on.")
after(h59._element, intro59._element)
rows59 = [
    ["Execution lag of two trading days (5.1)", "A signal needs the close of day s, so it cannot earn the return of day s; the position is traded on day s+1 and earns from day s+2, which avoids a same-bar look-ahead.", "Fixed for every sleeve"],
    ["NGL continuous series on the sheet's second and third columns (4.3)", "The first column of the NGL sheets is an average-price series that does not reflect where positions would be traded (project convention).", "Applies to Momentum and returns; Carry and Value use the literal columns"],
    ["NGL Carry pairs F2-F4 and F2-F14 (5.6)", "F4 was originally chosen to avoid heating-season seasonality, which F2 brings back (10 October note).", "Provisional"],
    ["Carry: V2 only, sign-vote across tenor pairs (5.3)", "User decision of 10 October 2026 to simplify Carry. Compared with the previous V1+V2 composite on full-sample IR for 2011 to mid-2026: V2 higher in Metals and Energy, slightly lower in NGL, not supportive in Precious Metals. The note records it as a decision to simplify, not a proven improvement.", "Compared in-sample; not significance-tested; regime-dependent"],
    ["Momentum: three crossovers (5.2)", "Compared with the three pairs of two on IR (10 October note): the default has the highest IR in Energy and NGL, and Momentum is negative in every version in Metals and Precious Metals, so there is no case to change it.", "Default kept; not significance-tested"],
    ["Two tenor pairs for Precious Metals and Agriculture, long leg about 12 months out (5.6)", "User decision of 10 to 11 October 2026. Precious short pair F1-F2; Platinum and Palladium F5 used as is (data from February 2010). The single-pair comparison is reported with the results.", "In-sample sensitivity study (ENGINE_CHANGELOG, 2026-10-11)"],
    ["Precious Metals and Agriculture calendars aligned to Bloomberg contract months (3.3)", "Exact price matching showed that Bloomberg's generic series switch contract only on each product's own cycle; the unaligned calendar produced spurious roll events (phase changes per Precious Metals product fell from 987 to between 331 and 495 after alignment).", "Data-quality fix"],
    ["Roll on the 5th trading day of the month of the last trade date (4.2)", "User decision of 10 to 11 October 2026. The Streamlit option 'Nth trading day of the expiry month' implements the same rule.", "Same rule for every product"],
    ["Agriculture limited to six products (2)", "Decision of 11 October 2026 (README); the six use the Bloomberg-matched calendar. The reason for excluding the other nine is not recorded in the repository.", "Reason to be confirmed"],
    ["Energy and the combined books shown with and without Natural Gas (2)", "Reported both ways so that the effect of the product is visible; no result is dropped.", "Presentation"],
    ["Value needs 630 observations of F3 (5.5); 504 in the 2012 variant (3.4)", "504 makes NGL Value and the NGL Inverse Vol series defined from 1 January 2012.", "2012 variant only"],
    ["Skewness: time-series sign rule on 252 observations (5.8)", "Same structure as the other sleeves; one year of daily returns as in Table 1 of the 2021 paper; the paper's cross-sectional sort is not used.", "Window taken from the paper; the Streamlit tab shows other windows for reference only"],
    ["Inverse Vol: 63 observations, lagged one day (9.2)", "Day t's weights use returns up to t-1, so the weights contain no look-ahead (lag added 11 October 2026).", "Fixed"],
    ["Costs: 5 bp round trip plus the roll-day charge (6)", "Convention used in every report in this project; the roll-day charge is half a round trip.", "A full 5 bp roll-day sensitivity has not been run"],
]
t59 = mk_table(["Decision", "Recorded basis", "Status"], rows59, [1.7, 3.5, 1.5])
after(intro59._element, t59._tbl)
tail59 = mk_par("Sensitivity studies on the same sample, recorded as local in ENGINE_CHANGELOG (2026-10-11): signal inversion (Metals, Precious Metals, Agriculture), Value look-back of 1, 3 and 5 years, combined-value versus sign-vote Carry and Carry-Momentum, the second tenor pair for Precious Metals and Agriculture, and a Platinum and Palladium F5 gap check. "
                "They informed the decisions above and are not reproduced in this document.")
after(t59._tbl, tail59._element)
# spacing paragraph between table and following heading is provided by tail59

# ---- Section 12: 2012 variant bullet
last12 = find("Rolling: Rolling Pearson correlation")
v12 = mk_par(" In the 2012 variant (Analysis/Final_Research_Dashboard_2012.html) all correlations start on 1 January 2012 and the NGL Inverse Vol series is defined from the start of the window, so the NGL date note above does not apply.", style="List Bullet", lead="2012 variant.")
after(last12._element, v12._element)

# ---- Section 14: rewrite Statistical status, add two bullets
ss = find("Statistical status.")
for r in list(ss.runs)[1:]:
    r._element.getparent().remove(r._element)
ss.add_run(" Results are a historical backtest of fixed rules over 2011 to 2026 (2012 to 2026 in the 2012 variant). No parameter is estimated or optimized by the code, and the regime and period splits are a check on stability, not an out-of-sample test. "
           "Several design decisions (Section 5.9) rest on comparisons run on the same sample, and no significance test, out-of-sample test or multiple-testing correction has been applied. Returns exclude financing, collateral income and any cost beyond the 5 bp round-trip convention.")
sel = mk_par(" Each of the 23 products, five sleeves, two constructions and ten pairs is a separate result inspected on the same data. The finding that the risk premium is concentrated in NGL and Energy was identified in this sample; with no correction for the number of comparisons it is in-sample evidence.",
             style="List Bullet", lead="Selection across many results.")
after(ss._element, sel._element)
nm = mk_par(" As Section 5.8(d) explains, the sleeve can carry directional exposure for long periods, so its correlations with equities, gold or a passive commodity index depend on the sample and should not be read as a stable property.",
            style="List Bullet", lead="Skewness is not market neutral.")
after(sel._element, nm._element)

# ---- Section 15: addition
s15 = find("The dashboard is produced by the scripts in scripts/final_dashboard")
add15 = mk_par("Added 12 October 2026. The Skewness sleeve is computed in compute_products.py and, separately, in scripts/final_dashboard/skewness_research.py (a research page that does not change the dashboards). "
               "The dashboard tables (asset-class tabs, all 23 products, both combined books, By Style and all Cross-Pair portfolios, 3,504 numbers) were recomputed from the per-product net returns with separate code and agree with the dashboard to within 0.0005 (Analysis/_v4_build/verify_independent.py). "
               "The Streamlit implementation is described in Section 16.")
after(s15._element, add15._element)

# ---- Section 16
h16 = mk_par("16. Streamlit Dashboards and the Hub", style="Heading 2")
after(add15._element, h16._element)
p16 = mk_par("The five Streamlit asset-class dashboards, their Portfolio tabs and the hub's Cross-Asset tab implement the same sleeves interactively. They are not the source of the Final Research Dashboard numbers and are not reconciled sleeve by sleeve with it; Carry in particular differs, as described below.")
after(h16._element, p16._element)
items16 = [
    ("Skewness tab.", " Each dashboard has a Skewness tab built like the Momentum tab (common_engine.render_skewness_tab): look-back windows of 126, 252 and 504 observations plus a custom one, with 252 as the research window, the default execution lag of Section 5.1 and the 5 bp cost. "
                      "The signal comes from one shared function, research/engine.py::raw_signal_skewness_from_returns, which the Portfolio tab and the hub also call."),
    ("Return convention.", " The standalone tabs earn dollar PnL on an additively stitched F1_continuous (a running sum of dollar changes), so the logarithm of that level is not a return: it divides by the adjusted level instead of the price of the contract held and is undefined where the level is not positive. "
                           "For Skewness the dashboards rebuild the exact ratio-stitched log return, log(1 + dollar change / price of the contract held at t-1), using F1_raw(t-1) on F1_Tracking days and F2_raw(t-1) on every other day (research/engine.py::ratio_log_returns). "
                           "This reproduces the returns of the series in Section 4.1 to about 1e-15 for all 23 products, and the signals agree on every day."),
    ("Portfolio tab.", " Skewness is a fifth strategy family (window 252, lag 2 days). The row 'Portfolio (all reference strategies)' keeps its four strategies and a new row adds Skewness. The tab uses the research log-price series, and its asset-class Skewness sleeve agrees with the Skewness sleeve of the 2012 variant dashboard for all five asset classes. "
                       "Its Carry reference is the V1 plus V2 composite per tenor pair, not the sign-vote V2 of Section 5.3, so its portfolio rows are not comparable with EW PORT."),
    ("Hub Cross-Asset tab.", " research/cross_asset_engine.py offers Skewness as an optional fifth style, off by default so that the four-style books are unchanged, and the hub adds an Inverse Vol combine method (1 / trailing 63-observation volatility, lagged one day) beside Equal Weight, Risk Parity and Dynamic Risk Parity. "
                             "Carry follows the 10 October rule and, with default settings, each product uses its own tenor pairs as in Section 5.6. With Skewness selected, the hub's five-style Equal Weight and Inverse Vol books agree with the By Style books of the 2012 variant dashboard to two decimals in IR."),
    ("Data basis.", " The hub and the dashboards read the repository configuration (Value needs 630 observations); the 2012 variant uses 504, which affects only the NGL products before 2012."),
]
prev = p16
for lead, text in items16:
    b = mk_par(text, style="List Bullet", lead=lead)
    after(prev._element, b._element)
    prev = b

tmp = PATH + ".tmp"
d.save(tmp)
os.replace(tmp, PATH)
print("methodology updated:", PATH)
