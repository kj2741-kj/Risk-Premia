"""Formula-driven tradebook: Brent Skewness sleeve (2026-10-12) -> tradebooks/Energy/Brent_Skewness_Tradebook.xlsx
Same layout and conventions as research/generate_tradebooks_logret.py and the earlier Wheat/Soybean tradebook: live Excel formulas from the price columns onward,
Python reference metrics next to the live-formula metrics, plus a static trade log. Data, roll rule (5th trading day of the expiry month) and the Brent data fix as in
the Final dashboard. Skewness sleeve: signal = minus the skewness of the daily log returns of the continuous series over the past 252 observations (Excel SKEW is the
same bias-adjusted sample skewness as pandas), position = signal lagged 2 rows, 5 bps costs with the roll-day charge. Run from the repo root."""
import os, sys, importlib, pickle
from pathlib import Path

os.environ["TRADEBOOK_ASSET_CLASS"] = "precious"   # only so the generator module imports (its output dir already exists)
R = Path(".").resolve()
for p in (R / "research", R / "research/configs", R / "scripts"):
    sys.path.insert(0, str(p))
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import generate_tradebooks_logret as g
from engine import exec_shift
from run_regime_table import _metrics

TC, WINDOW, SHIFT_N = 5, 252, 1
en = importlib.import_module("energy")
OUT_DIR = R / "tradebooks" / "Energy"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "Brent_Skewness_Tradebook.xlsx"

f1 = en.load_f1_series_logret("Brent")
assert not f1[["F1_raw", "F2_raw"]].isna().any().any(), "Brent has missing front prices; Excel formulas cannot rebuild the series"


def sheet_log_price(f):
    """Log ratio-continuous price rebuilt exactly as the sheet's formulas do (F1_Tracking: F1 ratio; Bridge: F1[t]/F2[t-1]; otherwise F2 ratio)."""
    G = [f["F1_raw"].iloc[0]]
    for i in range(1, len(f)):
        ph = f["Phase"].iloc[i]
        D, Dp, E, Ep = f["F1_raw"].iloc[i], f["F1_raw"].iloc[i - 1], f["F2_raw"].iloc[i], f["F2_raw"].iloc[i - 1]
        G.append(G[-1] * (D / Dp) if ph == "F1_Tracking" else G[-1] * (D / Ep) if ph == "Bridge" else G[-1] * (E / Ep))
    return pd.Series(np.log(G), index=f.index)


slp = sheet_log_price(f1)
pos = exec_shift(np.sign(-slp.diff().rolling(WINDOW).skew()), SHIFT_N).fillna(0)
ref, net = g._ref_metrics_from_position(pos, slp, f1["Phase"], TC)
ref = dict(ref)
for lab, start in (("2011-2026", "2011-01-01"), ("2012-2026", "2012-01-01")):
    m = _metrics(net.loc[start:])
    ref[f"Dashboard basis {lab}: Return (%)"] = round(m["ret"], 4)
    ref[f"Dashboard basis {lab}: Vol (%)"] = round(m["vol"], 4)
    ref[f"Dashboard basis {lab}: Max drawdown (%)"] = round(m["mdd"], 4)
    ref[f"Dashboard basis {lab}: IR"] = round(m["ir"], 4)

# reconcile with the dashboard's own Brent Skewness series
nets, _ = pickle.load(open("outputs/final_dashboard/products_v2_nets.pkl", "rb"))
dash = nets[("energy", "Brent")]["Skewness"]
diff = float((dash - net.reindex(dash.index)).abs().max())
print("max abs difference vs dashboard Brent Skewness series:", diff)

wb = Workbook()
readme = wb.active
readme.title = "README"
ws = wb.create_sheet("Brent_Skewness")
n_rows = len(f1)
eff_shift = SHIFT_N + 1
dates, f1_raw, f2_raw, phase = f1.index.tolist(), f1["F1_raw"].tolist(), f1["F2_raw"].tolist(), f1["Phase"].tolist()
d0 = g._data_start_row(len(ref))
d1 = d0 + n_rows - 1
live = {
    "Return (annualized %)": f"=AVERAGE(P{d0}:P{d1})*252*100",
    "Vol (annualized %)": f"=STDEV(P{d0}:P{d1})*SQRT(252)*100",
    "IR (Sharpe)": f"=AVERAGE(P{d0}:P{d1})/STDEV(P{d0}:P{d1})*SQRT(252)",
    "N active days": f"=COUNT(P{d0}:P{d1})",
}
g._perf_summary_block(ws, ref, live, n_cols=16)
headers = ["Date", "F1_raw", "F2_raw", "Phase", "F1_cont_ratio", "LogPrice", "LogReturn", "Skewness_252", "Skew_Signal",
           "Position", "Gross_PnL", "Position_Change", "TC_Cost", "Net_PnL", "Cum_Net_PnL", "Active_NetPnL"]


def row_fn(i):
    r = d0 + i
    v = [dates[i], f1_raw[i], f2_raw[i], phase[i], g._f1_cont_ratio_formula(r, i), f'=IF(E{r}<=0,"",LN(E{r}))']
    v.append("" if i == 0 else f'=IF(OR(F{r}="",F{r-1}=""),"",F{r}-F{r-1})')                       # G LogReturn
    v.append(f"=SKEW(G{r - WINDOW + 1}:G{r})" if i >= WINDOW else "")                             # H skewness of the last 252 log returns
    v.append(f'=IF(H{r}="","",-SIGN(H{r}))')                                                      # I signal: +1 if negative skew
    v.append(g._position_formula("I", r, i, eff_shift))                                           # J position (lagged 2 rows)
    v.append(f'=IF(OR(J{r}="",G{r}=""),0,J{r}*G{r})')                                             # K gross
    v.append(f"=ABS(J{r})" if i == 0 else f"=ABS(J{r}-J{r-1})")                                   # L position change
    v.append(g._tc_cost_formula("L", "J", "D", r, i, TC))                                         # M cost
    v.append(f"=K{r}-M{r}")                                                                       # N net
    v.append(f"=N{r}" if i == 0 else f"=O{r-1}+N{r}")                                             # O cumulative
    v.append(f'=IF(J{r}<>0,N{r},"")')                                                             # P active-day net
    return v


g._write_headers_and_rows(ws, headers, row_fn, n_rows, phase_col_idx=4)
for r in range(d0, ws.max_row + 1):
    if hasattr(ws.cell(r, 1).value, "year"):
        ws.cell(r, 1).number_format = "yyyy-mm-dd"

# static trade log
tl = wb.create_sheet("Brent_Trades")
tl.append(["Brent Skewness: position changes (static, from the Python engine; the daily sheet recomputes the same positions with formulas)"])
tl.cell(1, 1).font = Font(bold=True)
tl.append(["Date", "Position before", "Position after", "Action", "F1_raw (price)", "Contract held", "Phase", "Skewness_252 at signal"])
for c in range(1, 9):
    tl.cell(2, c).fill = PatternFill("solid", fgColor="2E5C8A"); tl.cell(2, c).font = Font(bold=True, color="FFFFFF", size=9)
p_ = pos.reindex(f1.index).fillna(0.0); prev = p_.shift(1).fillna(0.0)
sk = slp.diff().rolling(WINDOW).skew()
names = {1.0: "Long", -1.0: "Short", 0.0: "Flat"}
chg = p_[p_ != prev]
for d, p in chg.items():
    b = prev.loc[d]
    act = ("Open " + names[p]) if b == 0 else ("Close " + names[b]) if p == 0 else f"Flip {names[b]} to {names[p]}"
    i = f1.index.get_loc(d)
    s_ = sk.iloc[i - eff_shift] if i - eff_shift >= 0 else np.nan
    tl.append([d.to_pydatetime(), names[b], names[p], act, float(f1.at[d, "F1_raw"]), str(f1.at[d, "active_contract"]), str(f1.at[d, "Phase"]),
               None if pd.isna(s_) else round(float(s_), 4)])
    tl.cell(tl.max_row, 1).number_format = "yyyy-mm-dd"
for c, w in zip("ABCDEFGH", (12, 16, 16, 22, 16, 16, 14, 22)):
    tl.column_dimensions[c].width = w
tl.freeze_panes = "A3"

lines = [
    ["Brent Skewness tradebook (Skewness sleeve of the Final Research Dashboard)"], [""], ["Sheets"],
    ["Brent_Skewness", "Daily book, live Excel formulas from the price columns onward."],
    ["Brent_Trades", f"Static log of every position change ({len(chg)} changes), with the skewness value that drove each signal."], [""],
    ["Strategy definition"],
    ["Skewness", "LogReturn = daily change of LogPrice (the ratio-adjusted continuous Brent price). Skewness_252 = SKEW of the last 252 LogReturns (Excel SKEW is the "
                 "bias-adjusted sample skewness; only its sign is used). Skew_Signal = -SIGN(Skewness_252): +1 (long) when skewness is negative, -1 (short) when positive. "
                 "Position = signal lagged 2 rows (1-day execution lag plus the 1-day floor). First signal after 252 returns."],
    [""], ["Conventions"],
    ["Returns", "Log returns on F1_cont_ratio, rebuilt from Phase / F1_raw / F2_raw exactly as the engine does."],
    ["Costs", "5 bps round trip: half (2.5 bp) per unit of position change; on Roll_LTD days a held position pays the same half charge."],
    ["Roll", "5th trading day of the month in which each contract's last trade date falls (expiry_calendars_20260701.xlsx). "
             "Phase colours: orange = roll day, yellow = tracking F2, blue = day after expiry (bridge)."],
    ["Data", f"Bloomberg generic series, Energy_Futures_Updated.xlsx (Brent Crude ICE), {f1.index.min().date()} to {f1.index.max().date()}, {len(f1)} rows. "
             "The 2017-01-02 holiday row (empty curve) is removed, as in the dashboard."],
    ["Reading the summary", "Each sheet shows Python reference metrics next to the same metrics computed live by Excel; they should agree. Return/Vol/IR there use active days only "
                            "(days with a position). The 'Dashboard basis' rows are the full-span true-percent figures, as in the dashboards."],
    ["Reconciliation", f"Maximum absolute daily difference between this tradebook's Python net series and the dashboard's Brent Skewness series: {diff:.2e}."],
]
for ln in lines:
    readme.append(ln)
for r in (1, 3, 7, 10):
    readme.cell(r, 1).font = Font(bold=True)
readme.column_dimensions["A"].width = 26; readme.column_dimensions["B"].width = 140
for row in readme.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(OUT)
print("saved", OUT)
print("reference:", ref)
print("position changes:", len(chg), "rows:", n_rows)
